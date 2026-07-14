#!/usr/bin/env python3
"""Ensamblador mecánico de contratos ContractPerfectionist.

Este script NO redacta cláusulas. El contenido legal (objeto, alcance,
valores, plazos, cláusulas de protección) lo sigue redactando el
agente en conversación con el usuario, siguiendo las reglas de
CLAUDE.md (nunca inventar datos, consultar ante ambigüedad, etc.).

Lo único que resuelve este script es la parte mecánica que antes se
rehacía a mano en cada contrato:

  1. Construir el .docx SIEMPRE sobre `recursos/Contrato - Plantilla
     Base.docx`, preservando el membrete corporativo (regla 6).
  2. Insertar títulos, párrafos, tablas y el bloque de firmas con
     línea real (borde inferior de párrafo, regla 7) a partir de un
     JSON de datos ya confirmados.
  3. Verificar, antes de considerar el contrato listo para entrega,
     que no queden placeholders `[PENDIENTE ...]` sin resolver y que
     el documento conserve el membrete (regla 4 + gate de entrega).

Uso:
    python scripts/build_contract.py build --data contratos/<cliente>/datos.json --out "contratos/<cliente>/Contrato - <Cliente> - <YYYY-MM-DD>.docx"
    python scripts/build_contract.py check --file "contratos/<cliente>/Contrato - <Cliente> - <YYYY-MM-DD>.docx"

Ver scripts/README.md para el esquema completo del JSON de datos.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = REPO_ROOT / "recursos" / "Contrato - Plantilla Base.docx"
PLACEHOLDER_RE = re.compile(r"\[PENDIENTE[^\]]*\]", re.IGNORECASE)


# --------------------------------------------------------------------------
# Construcción
# --------------------------------------------------------------------------

def add_bottom_border(paragraph) -> None:
    """Añade un borde inferior real al párrafo (línea de firma física, regla 7)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_signature_block(doc: docx.document.Document, firmantes: list[dict]) -> None:
    """Agrega un bloque de firmas con línea física sobre cada nombre.

    Cada firmante: {"rol": "EL CONTRATISTA", "nombre": "...", "cargo": "...",
                     "identificacion": "..."} (identificacion es opcional).
    """
    doc.add_paragraph()
    table = doc.add_table(rows=0, cols=len(firmantes))
    table.autofit = True
    row_space = table.add_row()
    row_line = table.add_row()
    row_name = table.add_row()

    for i, firmante in enumerate(firmantes):
        row_space.cells[i].text = ""

        line_cell = row_line.cells[i]
        line_cell.text = ""
        p = line_cell.paragraphs[0]
        p.add_run(" ")  # espacio no separable para que el borde tenga longitud
        add_bottom_border(p)

        name_lines = [firmante["nombre"], firmante.get("rol", "")]
        if firmante.get("cargo"):
            name_lines.append(firmante["cargo"])
        if firmante.get("identificacion"):
            name_lines.append(firmante["identificacion"])

        name_cell = row_name.cells[i]
        name_cell.text = ""
        for j, line in enumerate(name_lines):
            para = name_cell.paragraphs[0] if j == 0 else name_cell.add_paragraph()
            run = para.add_run(line)
            run.font.size = Pt(11)
            if j == 0:
                run.bold = True
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table_from_data(doc: docx.document.Document, table_data: dict) -> None:
    """table_data: {"encabezados": [...], "filas": [[...], ...]}"""
    headers = table_data["encabezados"]
    rows = table_data["filas"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            row.cells[i].text = str(value)


def build_document(data: dict) -> docx.document.Document:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"No se encontró la plantilla corporativa en {TEMPLATE_PATH}. "
            "Nunca se genera un contrato sobre un documento en blanco (regla 6)."
        )

    doc = docx.Document(TEMPLATE_PATH)

    # El primer párrafo del template viene vacío: se usa como título.
    title_paragraph = doc.paragraphs[0]
    title_paragraph.style = doc.styles["Title"]
    title_paragraph.text = data["titulo"]
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for seccion in data.get("secciones", []):
        doc.add_heading(seccion["encabezado"], level=seccion.get("nivel", 1))
        for parrafo in seccion.get("parrafos", []):
            doc.add_paragraph(parrafo)
        if seccion.get("tabla"):
            add_table_from_data(doc, seccion["tabla"])

    if data.get("firmantes"):
        add_signature_block(doc, data["firmantes"])

    return doc


# --------------------------------------------------------------------------
# Gates de calidad (regla 1, regla 6, regla 7)
# --------------------------------------------------------------------------

def scan_placeholders(path: Path) -> list[str]:
    """Devuelve toda ocurrencia de `[PENDIENTE ...]` en cuerpo, tablas y encabezados."""
    doc = docx.Document(path)
    hits: list[str] = []

    def scan_paragraphs(paragraphs):
        for p in paragraphs:
            for m in PLACEHOLDER_RE.finditer(p.text):
                hits.append(m.group())

    scan_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                scan_paragraphs(cell.paragraphs)
    for section in doc.sections:
        scan_paragraphs(section.header.paragraphs)
        scan_paragraphs(section.footer.paragraphs)

    return hits


def verify_structure(path: Path) -> dict:
    """Verifica que el .docx conserve el membrete corporativo (regla 6) y
    que exista al menos un bloque de firmas con borde inferior (regla 7)."""
    result = {"header_image_ok": False, "signature_border_ok": False, "errors": []}

    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            header_files = [n for n in names if re.match(r"word/header\d*\.xml", n)]
            if not header_files:
                result["errors"].append("El documento no tiene ningún header de sección.")
            else:
                has_image_ref = False
                for hf in header_files:
                    rels_name = f"word/_rels/{Path(hf).name}.rels"
                    if rels_name in names:
                        rels_xml = z.read(rels_name).decode("utf-8", errors="ignore")
                        if "relationships/image" in rels_xml:
                            has_image_ref = True
                result["header_image_ok"] = has_image_ref
                if not has_image_ref:
                    result["errors"].append(
                        "Ningún header referencia una imagen: el membrete de Campuslands "
                        "pudo haberse perdido (regla 6)."
                    )

            doc_xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
            result["signature_border_ok"] = "<w:pBdr>" in doc_xml and "<w:bottom " in doc_xml
            if not result["signature_border_ok"]:
                result["errors"].append(
                    "No se encontró ningún párrafo con borde inferior: el bloque de firmas "
                    "puede no tener línea física (regla 7)."
                )
    except (KeyError, zipfile.BadZipFile) as exc:
        result["errors"].append(f"No se pudo inspeccionar el .docx: {exc}")

    return result


def run_check(path: Path) -> bool:
    """Corre ambos gates y reporta. Devuelve True si el contrato pasa limpio."""
    ok = True

    placeholders = scan_placeholders(path)
    if placeholders:
        ok = False
        print(f"[FALLO] Placeholders sin resolver ({len(placeholders)}):")
        for hit in placeholders:
            print(f"  - {hit}")
    else:
        print("[OK] Sin placeholders `[PENDIENTE ...]` pendientes.")

    structure = verify_structure(path)
    if structure["errors"]:
        ok = False
        print("[FALLO] Verificación de diseño corporativo:")
        for err in structure["errors"]:
            print(f"  - {err}")
    else:
        print("[OK] Membrete corporativo y línea de firma verificados.")

    return ok


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)

    build_p = sub.add_parser("build", help="Ensambla un contrato a partir de un JSON de datos.")
    build_p.add_argument("--data", required=True, type=Path, help="Ruta al JSON de datos (ver scripts/README.md).")
    build_p.add_argument("--out", required=True, type=Path, help="Ruta de salida del .docx.")
    build_p.add_argument("--skip-check", action="store_true", help="No correr los gates de calidad al terminar.")

    check_p = sub.add_parser("check", help="Corre los gates de calidad sobre un .docx ya existente.")
    check_p.add_argument("--file", required=True, type=Path)

    args = parser.parse_args()

    if args.command == "build":
        data = json.loads(args.data.read_text(encoding="utf-8"))
        doc = build_document(data)
        args.out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(args.out)
        print(f"Contrato generado en: {args.out}")
        if not args.skip_check:
            print()
            passed = run_check(args.out)
            if not passed:
                print(
                    "\nEl contrato NO debe moverse a estado de entrega hasta resolver lo anterior."
                )
                return 1
        return 0

    if args.command == "check":
        return 0 if run_check(args.file) else 1

    return 1


if __name__ == "__main__":
    sys.exit(main())
